from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / "docs" / "ACCEPTANCE_TEST_PLAN_VERSION2.zh-CN.md"
target = ROOT / "Hospital_Patient_Pathway_Acceptance_Test_Plan.docx"
lines = source.read_text(encoding="utf-8").splitlines()
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.75)
section.right_margin = Inches(.75)

styles = doc.styles
for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Aptos"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(4)
styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)

def add_inline(p, text):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]*`|\[[^\]]+\]\([^)]*\))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(8)
        elif part.startswith("["):
            p.add_run(re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", part))
        else:
            p.add_run(part)

i = 0
while i < len(lines):
    s = lines[i].strip()
    if not s:
        i += 1; continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            row = [x.strip() for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", x) for x in row):
                rows.append(row)
            i += 1
        if rows:
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            for j, val in enumerate(rows[0]):
                cell = table.rows[0].cells[j]; cell.text = val; set_cell_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs: run.bold = True; run.font.size = Pt(8)
            for row in rows[1:]:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    cells[j].text = val
                    for p in cells[j].paragraphs:
                        p.paragraph_format.space_after = Pt(1)
                        for run in p.runs: run.font.size = Pt(8)
            doc.add_paragraph()
        continue
    if s.startswith("``` ") or s == "```":
        i += 1; block = []
        while i < len(lines) and lines[i].strip() != "```": block.append(lines[i]); i += 1
        p = doc.add_paragraph()
        run = p.add_run("\n".join(block)); run.font.name = "Consolas"; run.font.size = Pt(8)
        i += 1; continue
    m = re.match(r"^(#{1,3})\s+(.*)$", s)
    if m:
        level = len(m.group(1)); heading = doc.add_heading("", level=level if level < 4 else 3); add_inline(heading, m.group(2))
    elif s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); add_inline(p, s[2:])
    else:
        p = doc.add_paragraph(); add_inline(p, s)
    i += 1

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Hospital Patient Pathway · Version 2 Acceptance Test Plan")

doc.core_properties.title = "Hospital Patient Pathway Version 2 Acceptance Test Plan"
doc.core_properties.subject = "Version 1 acceptance criteria mapped to Version 2"
doc.save(target)
print(target)
